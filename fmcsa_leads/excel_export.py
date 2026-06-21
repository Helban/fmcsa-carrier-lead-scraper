"""Export carrier leads to an Excel workbook with a QA summary sheet."""
from __future__ import annotations

import logging
from collections import Counter

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .leads import CarrierLead

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="217346")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BAND_FILL = PatternFill(fill_type="solid", fgColor="EDF3EF")
_WIDTH_PADDING = 2  # breathing room past the longest cell in a column
_MAX_COLUMN_WIDTH = 42  # cap so one long cargo string can't blow out the layout


def export_leads(leads: list[CarrierLead], output_path: str) -> None:
    """Write a Leads sheet and a QA Summary sheet to ``output_path`` (.xlsx)."""
    lead_columns = {
        "DOT#": [lead.dot_number for lead in leads],
        "Company": [lead.legal_name for lead in leads],
        "Phone": [lead.phone for lead in leads],
        "Street": [lead.street for lead in leads],
        "City": [lead.city for lead in leads],
        "State": [lead.state for lead in leads],
        "ZIP": [lead.zip_code for lead in leads],
        "Power units": [lead.power_units for lead in leads],
        "Trucks": [lead.truck_units for lead in leads],
        "Drivers": [lead.total_drivers for lead in leads],
        "Cargo": [lead.cargo for lead in leads],
    }
    if _was_enriched(leads):
        lead_columns["Website"] = [lead.website for lead in leads]
        lead_columns["Email"] = [lead.email for lead in leads]
    leads_frame = pd.DataFrame(lead_columns)
    qa_frame = _build_qa_summary(leads)
    with pd.ExcelWriter(output_path, engine="openpyxl") as excel_writer:
        leads_frame.to_excel(excel_writer, sheet_name="Leads", index=False)
        qa_frame.to_excel(excel_writer, sheet_name="QA Summary", index=False)
        leads_worksheet = excel_writer.sheets["Leads"]
        _style_sheet(leads_worksheet, leads_frame)
        _band_alternate_rows(leads_worksheet, leads_frame)
        _enable_header_filter(leads_worksheet)
        _style_sheet(excel_writer.sheets["QA Summary"], qa_frame)
    logger.info("Wrote %d leads to %s", len(leads), output_path)


def _style_sheet(worksheet: Worksheet, frame: pd.DataFrame) -> None:
    """Green bold header, frozen top row, and content-fit column widths."""
    for column_position, column_name in enumerate(frame.columns, start=1):
        header_cell = worksheet.cell(row=1, column=column_position)
        header_cell.fill = _HEADER_FILL
        header_cell.font = _HEADER_FONT
        longest_value = max(
            (len(str(value)) for value in frame.iloc[:, column_position - 1]),
            default=0,
        )
        content_width = max(len(str(column_name)), longest_value) + _WIDTH_PADDING
        worksheet.column_dimensions[get_column_letter(column_position)].width = min(
            content_width, _MAX_COLUMN_WIDTH
        )
    worksheet.freeze_panes = "A2"


def _band_alternate_rows(worksheet: Worksheet, frame: pd.DataFrame) -> None:
    """Shade every second data row so a long list stays readable."""
    column_count = len(frame.columns)
    for data_position, sheet_row in enumerate(range(2, 2 + len(frame)), start=1):
        if data_position % 2 == 0:
            for column_position in range(1, column_count + 1):
                worksheet.cell(row=sheet_row, column=column_position).fill = _BAND_FILL


def _enable_header_filter(worksheet: Worksheet) -> None:
    """Turn the header row into Excel dropdown filters."""
    worksheet.auto_filter.ref = worksheet.dimensions


def _was_enriched(leads: list[CarrierLead]) -> bool:
    return any(lead.website or lead.email for lead in leads)


def _build_qa_summary(leads: list[CarrierLead]) -> pd.DataFrame:
    """Build a metric/value table the end user can sanity-check the list against."""
    lead_count = len(leads)
    with_phone = sum(1 for lead in leads if lead.phone)
    by_state = Counter(lead.state for lead in leads)
    phone_coverage = round(100 * with_phone / lead_count, 1) if lead_count else 0.0
    top_state, top_state_count = by_state.most_common(1)[0] if by_state else ("", 0)
    qa_metrics = [
        ("Total leads", lead_count),
        ("With phone", with_phone),
        ("Phone coverage %", phone_coverage),
        ("Distinct states", len(by_state)),
        ("Largest state", f"{top_state} ({top_state_count})"),
    ]
    if _was_enriched(leads):
        with_email = sum(1 for lead in leads if lead.email)
        email_coverage = round(100 * with_email / lead_count, 1) if lead_count else 0.0
        qa_metrics.append(("With email", with_email))
        qa_metrics.append(("Email coverage %", email_coverage))
    return pd.DataFrame(qa_metrics, columns=["Metric", "Value"])
