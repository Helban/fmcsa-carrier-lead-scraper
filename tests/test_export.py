"""Tests for the styled Excel export (no network)."""
from pathlib import Path

import openpyxl

from fmcsa_leads.excel_export import export_leads
from fmcsa_leads.leads import CarrierLead

_EXCEL_DEFAULT_COLUMN_WIDTH = 8.43


def _powder_river_lead() -> CarrierLead:
    return CarrierLead(
        dot_number="19446",
        legal_name="POWDER RIVER TRANSPORTATION",
        phone="(307) 487-0653",
        street="1700 E HIGHWAY 14/16",
        city="GILLETTE",
        state="WY",
        zip_code="82716",
        power_units=52,
        truck_units=50,
        total_drivers=48,
        cargo="General Freight",
    )


def test_export_writes_both_sheets(tmp_path: Path) -> None:
    output_path = tmp_path / "leads.xlsx"
    export_leads([_powder_river_lead()], str(output_path))
    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == ["Leads", "QA Summary"]


def test_header_is_filled_bold_and_frozen(tmp_path: Path) -> None:
    output_path = tmp_path / "leads.xlsx"
    export_leads([_powder_river_lead()], str(output_path))
    leads_sheet = openpyxl.load_workbook(output_path)["Leads"]
    header_cell = leads_sheet["A1"]
    assert header_cell.font.bold is True
    assert header_cell.fill.fgColor.rgb.endswith("217346")
    assert leads_sheet.freeze_panes == "A2"


def test_long_company_column_widened_past_excel_default(tmp_path: Path) -> None:
    output_path = tmp_path / "leads.xlsx"
    export_leads([_powder_river_lead()], str(output_path))
    leads_sheet = openpyxl.load_workbook(output_path)["Leads"]
    # Company (column B) holds a 27-char name, so its width must exceed the default.
    assert leads_sheet.column_dimensions["B"].width > _EXCEL_DEFAULT_COLUMN_WIDTH


def test_leads_sheet_has_autofilter(tmp_path: Path) -> None:
    output_path = tmp_path / "leads.xlsx"
    export_leads([_powder_river_lead()], str(output_path))
    leads_sheet = openpyxl.load_workbook(output_path)["Leads"]
    assert leads_sheet.auto_filter.ref is not None


def test_alternating_data_rows_are_banded(tmp_path: Path) -> None:
    output_path = tmp_path / "leads.xlsx"
    export_leads([_powder_river_lead()] * 3, str(output_path))
    leads_sheet = openpyxl.load_workbook(output_path)["Leads"]
    assert leads_sheet["A2"].fill.fill_type is None  # first data row stays plain
    assert leads_sheet["A3"].fill.fgColor.rgb.endswith("EDF3EF")  # second data row banded
